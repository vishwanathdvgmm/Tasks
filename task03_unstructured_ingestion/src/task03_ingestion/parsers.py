from __future__ import annotations

from datetime import datetime, timezone
from pathlib import Path
import re
from typing import Iterable

import pymupdf
from bs4 import BeautifulSoup
from docx import Document as DocxDocument

from .cleaning import clean_text
from .models import DocumentRecord

SUPPORTED = {".pdf", ".md", ".markdown", ".txt", ".docx", ".html", ".htm"}

def _timestamp(value: object) -> str | None:
    if value is None:
        return None
    value = str(value).strip()
    if not value:
        return None
    # PDF metadata often uses D:YYYYMMDDHHmmSS offsets; retain a safe normalized value.
    match = re.match(r"D:(\d{4})(\d{2})?(\d{2})?(\d{2})?(\d{2})?(\d{2})?", value)
    if match:
        parts = [p or "01" for p in match.groups()]
        try:
            return datetime(*[int(p) for p in parts]).replace(tzinfo=timezone.utc).isoformat()
        except ValueError:
            return value
    return value

def _record(path: Path, text: str, **kwargs) -> DocumentRecord:
    return DocumentRecord(
        text=clean_text(text),
        source_filename=path.name,
        source_path=str(path),
        **kwargs,
    )

def parse_file(path: Path, ocr: bool = False) -> list[DocumentRecord]:
    """Parse one supported file into one or more page/section records."""
    ext = path.suffix.lower()
    if ext not in SUPPORTED:
        raise ValueError(f"Unsupported file extension: {ext}")

    if ext == ".pdf":
        records: list[DocumentRecord] = []
        with pymupdf.open(path) as pdf:
            meta = pdf.metadata or {}
            author = meta.get("author") or None
            created = _timestamp(meta.get("creationDate"))
            for index, page in enumerate(pdf, start=1):
                text = page.get_text("text")
                if len(text.strip()) < 20 and ocr:
                    try:
                        import importlib
                        import io

                        pytesseract = importlib.import_module("pytesseract")
                        pil_image = importlib.import_module("PIL.Image")
                        pix = page.get_pixmap(matrix=pymupdf.Matrix(2, 2), alpha=False)
                        image = pil_image.open(io.BytesIO(pix.tobytes("png")))
                        text = pytesseract.image_to_string(image)
                    except ImportError as exc:
                        raise RuntimeError(
                            "OCR requires `uv sync --extra ocr` and a system Tesseract install."
                        ) from exc
                # Best-effort table extraction; availability depends on PyMuPDF version/document.
                try:
                    tables = page.find_tables()
                    table_blocks = []
                    for table in tables.tables:
                        rows = table.extract()
                        if rows:
                            table_blocks.append(_table_to_markdown(rows))
                    if table_blocks:
                        text += "\n\n" + "\n\n".join(table_blocks)
                except (AttributeError, RuntimeError, ValueError):
                    pass
                records.append(_record(
                    path, text, page_number=index, author=author,
                    document_created_at=created,
                    extra={"format": "pdf", "page_width": page.rect.width,
                           "page_height": page.rect.height,
                           "ocr_requested": ocr},
                ))
        return records

    if ext in {".md", ".markdown", ".txt"}:
        return [_record(path, path.read_text(encoding="utf-8", errors="replace"),
                        extra={"format": "markdown" if ext != ".txt" else "text"})]

    if ext == ".docx":
        doc = DocxDocument(path)
        parts: list[str] = []
        for paragraph in doc.paragraphs:
            if paragraph.text.strip():
                parts.append(paragraph.text)
        for table in doc.tables:
            rows = [[cell.text.strip().replace("\n", " ") for cell in row.cells]
                    for row in table.rows]
            parts.append(_table_to_markdown(rows))
        core = doc.core_properties
        created = core.created.isoformat() if core.created else None
        return [_record(path, "\n\n".join(parts), author=core.author or None,
                        document_created_at=created, extra={"format": "docx"})]

    if ext in {".html", ".htm"}:
        soup = BeautifulSoup(path.read_text(encoding="utf-8", errors="replace"), "html.parser")
        for tag in soup(["script", "style", "noscript", "svg"]):
            tag.decompose()
        parts: list[str] = []
        title = soup.title.get_text(" ", strip=True) if soup.title else None
        for table in soup.find_all("table"):
            rows = []
            for tr in table.find_all("tr"):
                cells = [cell.get_text(" ", strip=True) for cell in tr.find_all(["th", "td"])]
                if cells:
                    rows.append(cells)
            if rows:
                parts.append(_table_to_markdown(rows))
            table.decompose()
        body = soup.body or soup
        text = body.get_text("\n", strip=True)
        return [_record(path, (f"# {title}\n\n" if title else "") + text
                        + ("\n\n" + "\n\n".join(parts) if parts else ""),
                        section_header=title, extra={"format": "html"})]

    raise ValueError(f"Parser not implemented for {ext}")

def _table_to_markdown(rows: Iterable[Iterable[object]]) -> str:
    materialized = [[str(cell or "").replace("|", r"\|").replace("\n", " ").strip()
                     for cell in row] for row in rows]
    if not materialized:
        return ""
    width = max(len(row) for row in materialized)
    materialized = [row + [""] * (width - len(row)) for row in materialized]
    header = materialized[0]
    separator = ["---"] * width
    lines = ["| " + " | ".join(header) + " |",
             "| " + " | ".join(separator) + " |"]
    lines.extend("| " + " | ".join(row) + " |" for row in materialized[1:])
    return "\n".join(lines)

def discover_files(path: Path) -> list[Path]:
    if path.is_file():
        return [path] if path.suffix.lower() in SUPPORTED else []
    return sorted(p for p in path.rglob("*") if p.is_file() and p.suffix.lower() in SUPPORTED)
